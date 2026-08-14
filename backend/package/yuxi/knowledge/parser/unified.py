"""Unified parser module for markdown conversion."""

from __future__ import annotations

import asyncio
import base64
import os
import re
import tempfile
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import aiofiles
from docling.datamodel.base_models import InputFormat
from docling.document_converter import DocumentConverter
from langchain_community.document_loaders import PyPDFLoader
from markdownify import markdownify as md_convert

from yuxi.knowledge.parser.zip_utils import process_zip_file as _process_zip_file
from yuxi.storage.minio import get_minio_client
from yuxi.utils import logger

SUPPORTED_FILE_EXTENSIONS: tuple[str, ...] = (
    ".txt",
    ".md",
    ".docx",
    ".html",
    ".htm",
    ".json",
    ".csv",
    ".xls",
    ".xlsx",
    ".pdf",
    ".pptx",
    ".jpg",
    ".jpeg",
    ".png",
    ".bmp",
    ".tiff",
    ".tif",
    ".zip",
)


def is_supported_file_extension(file_name: str | os.PathLike[str]) -> bool:
    """Check whether the given file path has a supported extension."""
    return Path(file_name).suffix.lower() in SUPPORTED_FILE_EXTENSIONS


@dataclass(slots=True)
class MarkdownParseResult:
    """Unified Markdown parsing results."""

    markdown: str
    file_ext: str | None = None
    artifacts: dict[str, Any] = field(default_factory=dict)


_docling_converter: DocumentConverter | None = None


def _get_docling_converter() -> DocumentConverter:
    """Get the Docling document converter singleton."""
    global _docling_converter
    if _docling_converter is None:
        _docling_converter = DocumentConverter(
            format_options={
                InputFormat.DOCX: None,
                InputFormat.XLSX: None,
                InputFormat.PPTX: None,
            }
        )
    return _docling_converter


def _resolve_image_storage_params(params: dict | None) -> tuple[str, str]:
    params = params or {}

    image_bucket = params.get("image_bucket") or "knowledgebases"
    image_prefix = params.get("image_prefix")
    if image_prefix:
        normalized_prefix = str(image_prefix).strip("/")
        if normalized_prefix:
            return image_bucket, normalized_prefix

    return image_bucket, "unknown/kb-images"


def _resolve_ocr_engine_params(params: dict | None) -> tuple[str, dict[str, Any]]:
    from yuxi import config

    params = params or {}
    engine = str(params.get("ocr_engine") if "ocr_engine" in params else config.default_ocr_engine)
    engine = engine.strip() or config.default_ocr_engine
    engine_config = params.get("ocr_engine_config")
    processor_params = dict(params)
    if isinstance(engine_config, dict):
        processor_params.update(engine_config)
    return engine, processor_params


def _upload_image_to_minio(image_data: bytes, filename: str, bucket_name: str, object_prefix: str) -> str:
    """Upload images to MinIO and return URL."""
    minio_client = get_minio_client()
    minio_client.ensure_bucket_exists(bucket_name)

    normalized_prefix = object_prefix.strip("/") or "unknown/kb-images"
    timestamp = int(time.time() * 1000000)
    object_name = f"{normalized_prefix}/{timestamp}_{Path(filename).name}"

    result = minio_client.upload_file(
        bucket_name=bucket_name,
        object_name=object_name,
        data=image_data,
    )
    if bucket_name == minio_client.KB_BUCKETS["images"]:
        return f"/api/v1/knowledge/images/{object_name}"
    return result.url


def _parse_data_uri(data_uri: str) -> tuple[bytes, str]:
    """Parse the data URI and return (image_data, mime_type)。"""
    header, base64_data = data_uri.split(",", 1)
    mime_type = header.split(":")[1].split(";")[0]
    image_data = base64.b64decode(base64_data)
    return image_data, mime_type


def _convert_with_docling(file_path: Path, params: dict | None = None) -> str:
    """Use Docling to convert docx/xlsx/pptx Convert to Markdown."""
    params = params or {}
    image_bucket, image_prefix = _resolve_image_storage_params(params)

    converter = _get_docling_converter()
    result = converter.convert(file_path)

    if result.status.name != "SUCCESS":
        raise RuntimeError(f"Chuyển đổi Docling thất bại: {result.status}")

    doc = result.document

    if hasattr(doc, "pictures") and doc.pictures:
        replacements: list[str] = []
        for pic in doc.pictures:
            uri = str(pic.image.uri) if hasattr(pic, "image") and hasattr(pic.image, "uri") else ""
            if uri.startswith("data:"):
                filename = "image"
                try:
                    image_data, mime_type = _parse_data_uri(uri)
                    filename = f"image_{int(time.time() * 1000000)}.{mime_type.split('/')[-1]}"
                    url = _upload_image_to_minio(image_data, filename, image_bucket, image_prefix)
                    replacements.append(f"![{filename}]({url})")
                except Exception as e:  # noqa: BLE001
                    logger.error(f"Failed to upload image {filename}: {e}")
                    replacements.append(f"[picture: {filename}]")
            else:
                replacements.append("")

        markdown = doc.export_to_markdown()
        for replacement in replacements:
            markdown = re.sub(r"<!--\s*image\s*-->", replacement, markdown, count=1)
        return markdown

    return doc.export_to_markdown()


def _convert_docx_with_python_docx(file_path: Path) -> str:
    """using python-docx Parse DOCX (cover when Docling fails)."""
    from docx import Document

    document = Document(str(file_path))
    blocks: list[str] = []

    for para in document.paragraphs:
        text = para.text.strip()
        if text:
            blocks.append(text)

    for table in document.tables:
        rows: list[list[str]] = []
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                rows.append(cells)

        if not rows:
            continue

        header = rows[0]
        blocks.append(f"| {' | '.join(header)} |")
        blocks.append(f"| {' | '.join(['---'] * len(header))} |")

        for row in rows[1:]:
            normalized_row = row + [""] * (len(header) - len(row))
            blocks.append(f"| {' | '.join(normalized_row[: len(header)])} |")

        blocks.append("")

    return "\n\n".join(blocks).strip()


def _convert_csv_to_markdown(file_path: Path) -> str:
    import pandas as pd

    dataframe = pd.read_csv(file_path)
    tables: list[str] = []
    for i in range(len(dataframe)):
        row_dataframe = dataframe.iloc[[i]]
        tables.append(row_dataframe.to_markdown(index=False))
    return "\n\n".join(tables)


def pdfreader(file_path, params=None):
    """Reads a PDF file and returns text."""
    if isinstance(file_path, str):
        file_path = Path(file_path)

    assert file_path.exists(), "File not found"
    assert file_path.suffix.lower() == ".pdf", "File format not supported"

    loader = PyPDFLoader(str(file_path))
    docs = loader.load()
    text = "\n\n".join([d.page_content for d in docs])
    return text


def parse_pdf(file, params=None):
    """Parse PDF files and support multiple OCR methods."""
    from yuxi.knowledge.parser.base import DocumentProcessorException
    from yuxi.knowledge.parser.factory import DocumentProcessorFactory
    from yuxi.knowledge.parser.density import PDFDensityAnalyzer
    from yuxi.knowledge.parser.models import OCRPolicy, ProcessingStatus, ProcessingResult

    from yuxi import config as yuxi_config

    params = params or {}
    trace_id = params.get("file_id") or Path(str(file)).name

    # 0. Preflight check for PDF page tree loadability - ponytail: validate PDF page tree before parsing
    from yuxi.knowledge.utils.pdf_utils import validate_pdf_page_tree_loadable
    validate_pdf_page_tree_loadable(file)

    # 1. Analyzer
    try:
        analyzer = PDFDensityAnalyzer()
        analysis = analyzer.analyze_document(file)
        logger.info(
            f"[FileID: {trace_id}] [Density Analyzer] Recommended: {analysis['recommended_ocr']}, "
            f"Scan pages: {len(analysis['scan_pages'])}/{analysis['total_pages']}, "
            f"Hybrid pages: {len(analysis['hybrid_pages'])}/{analysis['total_pages']}"
        )
    except Exception as e:
        logger.warning(f"[FileID: {trace_id}] [Density Analyzer] Failed to analyze PDF density: {e}")
        analysis = {"recommended_ocr": "auto"}

    # Policy: explicit param wins, then runtime config, then density-based auto.
    # An explicit ocr_engine="disable" means "text-layer only, no OCR" and always forces DISABLE.
    opt_ocr, _ = _resolve_ocr_engine_params(params)
    policy_str = params.get("ocr_policy") or yuxi_config.ocr_policy or "auto"
    try:
        policy = OCRPolicy(str(policy_str))
    except ValueError:
        policy = OCRPolicy.AUTO

    if opt_ocr == "disable":
        policy = OCRPolicy.DISABLE
    elif policy == OCRPolicy.AUTO:
        policy = OCRPolicy.DISABLE if analysis.get("recommended_ocr") == "disable" else OCRPolicy.ENABLE

    allow_external = yuxi_config.allow_external_ocr

    # Engine chain: user-selected engine first, then registered defaults.
    engine_preference: list[str] = []
    
    if opt_ocr and opt_ocr != "disable" and opt_ocr in DocumentProcessorFactory.PROCESSOR_TYPES:
        if DocumentProcessorFactory.requires_external(opt_ocr) and not allow_external:
            logger.error(f"[FileID: {trace_id}] Compliance violation: attempt to use external OCR '{opt_ocr}' when ALLOW_EXTERNAL_OCR=false.")
            from yuxi.knowledge.parser.base import DocumentProcessorException
            raise DocumentProcessorException(
                f"Cannot use {opt_ocr}: External OCR is disabled by compliance policy (ALLOW_EXTERNAL_OCR=false).", 
                opt_ocr
            )
        engine_preference.append(opt_ocr)
    if policy == OCRPolicy.DISABLE:
        # Text-layer PDFs: local text parser without OCR first, then local OCR fallback.
        if "docling" not in engine_preference:
            engine_preference.append("docling")
        if "rapid_ocr" not in engine_preference:
            engine_preference.append("rapid_ocr")
    else:
        # Scanned/hybrid PDFs: prefer cloud OCR when allowed, otherwise local OCR.
        prioritized_engines = [
            "paddleocr_vl_1_6",
            "mineru_official",
            "deepseek_ocr",
            "paddleocr_pp_ocrv6",
            "rapid_ocr",
            "pp_structure_v3_ocr",
            "mineru_ocr"
        ]
        for engine in prioritized_engines:
            if engine not in DocumentProcessorFactory.PROCESSOR_TYPES:
                continue
            try:
                is_ext = DocumentProcessorFactory.requires_external(engine)
            except Exception:
                is_ext = False
            
            if is_ext and not allow_external:
                continue
            if engine not in engine_preference:
                engine_preference.append(engine)

    # Resolve images
    image_bucket, image_prefix = _resolve_image_storage_params(params)
    processor_params = dict(params)
    processor_params.setdefault("image_bucket", image_bucket)
    processor_params.setdefault("image_prefix", image_prefix)
    processor_params["ocr_policy"] = policy.value

    # Execution
    last_error = None
    for engine in engine_preference:
        try:
            processor = DocumentProcessorFactory.get_processor(engine)
            if hasattr(processor, "process"):
                result = processor.process(file, processor_params)
            else:
                # Backward-compatible wrapper for processors that only implement process_file
                content = processor.process_file(file, processor_params)
                result = ProcessingResult(
                    status=ProcessingStatus.SUCCESS,
                    engine=engine,
                    ocr_used=(policy != OCRPolicy.DISABLE),
                    content=content,
                )

            if result.status == ProcessingStatus.SUCCESS:
                logger.info(f"[FileID: {trace_id}] [Engine: {engine}] SUCCESS, content length={len(result.content or '')}")
                return result

            if result.status == ProcessingStatus.DEGRADED:
                logger.warning(f"[FileID: {trace_id}] [Engine: {engine}] DEGRADED result, trying next engine...")
                last_error = DocumentProcessorException(f"Degraded result from {engine}", engine)
                continue

        except Exception as e:
            last_error = DocumentProcessorException(str(e), engine)
            logger.warning(f"[FileID: {trace_id}] [Engine: {engine}] failed: {e}")

    # No automatic fallback to the plain-text PyPDF reader; surface the failure explicitly.
    return ProcessingResult(
        status=ProcessingStatus.FAILED,
        engine="orchestrator",
        ocr_used=False,
        error=last_error,
        metadata={"fallback_chain_exhausted": True},
    )


def parse_image(file, params=None):
    """Parse image files and support multiple OCR methods."""
    from yuxi.knowledge.parser.base import DocumentProcessorException
    from yuxi.knowledge.parser.factory import DocumentProcessorFactory

    opt_ocr, processor_params = _resolve_ocr_engine_params(params)

    if opt_ocr == "disable":
        raise ValueError(
            "Tệp hình ảnh bắt buộc phải bật OCR để trích xuất nội dung văn bản. "
            "Vui lòng chọn phương thức OCR (rapid_ocr/mineru_ocr/mineru_official/pp_structure_v3_ocr/deepseek_ocr/"
            "paddleocr_vl_1_6/paddleocr_pp_ocrv6) hoặc xóa tệp này."
        )

    image_bucket, image_prefix = _resolve_image_storage_params(processor_params)
    processor_params.setdefault("image_bucket", image_bucket)
    processor_params.setdefault("image_prefix", image_prefix)

    try:
        return DocumentProcessorFactory.process_file(opt_ocr, file, processor_params)
    except DocumentProcessorException as e:
        logger.error(f"Image processing failed: {e.service_name} - {str(e)}")
        raise
    except Exception as e:  # noqa: BLE001
        logger.error(f"Phân tích hình ảnh thất bại: {str(e)}")
        raise DocumentProcessorException(f"Phân tích hình ảnh thất bại: {str(e)}", opt_ocr, "parsing_failed")


async def parse_pdf_async(file, params=None):
    return await asyncio.to_thread(parse_pdf, file, params=params)


async def parse_image_async(file, params=None):
    return await asyncio.to_thread(parse_image, file, params=params)


async def _process_file_to_markdown_core(
    file_path: str, params: dict | None = None
) -> tuple[str, str | None, dict[str, Any]]:
    """Convert different types of files to markdown, supporting local files and MinIO files."""
    from yuxi.knowledge.utils.kb_utils import is_minio_url, parse_minio_url
    from yuxi.storage.minio.client import get_minio_client

    if is_minio_url(file_path):
        logger.debug(f"Downloading file from MinIO: {file_path}")

        if "?" in file_path:
            file_path_clean = file_path.split("?")[0]
        else:
            file_path_clean = file_path

        original_filename = file_path_clean.split("/")[-1]

        with tempfile.NamedTemporaryFile(delete=False, suffix=Path(original_filename).suffix) as temp_file:
            temp_path = temp_file.name

        try:
            bucket_name, object_name = parse_minio_url(file_path)
            minio_client = get_minio_client()
            file_content = await minio_client.adownload_file(bucket_name, object_name)

            async with aiofiles.open(temp_path, "wb") as f:
                await f.write(file_content)

            logger.debug(f"File downloaded to temp path: {temp_path}")
            actual_file_path = temp_path

        except Exception as e:  # noqa: BLE001
            if os.path.exists(temp_path):
                os.unlink(temp_path)
            logger.error(f"Failed to download file from MinIO: {e}")
            raise ValueError(f"Không thể tải tệp xuống từ MinIO: {e}")
    else:
        actual_file_path = file_path

    file_ext: str | None = None
    artifacts: dict[str, Any] = {}

    try:
        file_path_obj = Path(actual_file_path)
        file_ext = file_path_obj.suffix.lower()

        if file_ext == ".pdf":
            parse_result = await parse_pdf_async(str(file_path_obj), params=params)
            if not parse_result.success:
                from yuxi.knowledge.parser.base import DocumentProcessorException

                raise DocumentProcessorException(f"Parse PDF failed: {parse_result.error}", parse_result.engine)
            text = parse_result.content
            result = f"{text}"

        elif file_ext in [".txt", ".md"]:
            async with aiofiles.open(file_path_obj, encoding="utf-8") as f:
                content = await f.read()
            result = f"{content}"

        elif file_ext == ".docx":
            try:
                result = await asyncio.to_thread(_convert_with_docling, file_path_obj, params=params)
            except Exception as e:  # noqa: BLE001
                logger.warning(f"Docling Failed to parse DOCX, fallback to python-docx: {file_path_obj.name}, {e}")
                result = await asyncio.to_thread(_convert_docx_with_python_docx, file_path_obj)

        elif file_ext == ".pptx":
            result = await asyncio.to_thread(_convert_with_docling, file_path_obj, params=params)

        elif file_ext == ".doc":
            from langchain_community.document_loaders import UnstructuredWordDocumentLoader

            loader = UnstructuredWordDocumentLoader(str(file_path_obj))
            docs = await asyncio.to_thread(loader.load)
            result = "\n".join(doc.page_content for doc in docs).strip()

        elif file_ext in [".jpg", ".jpeg", ".png", ".bmp", ".tiff", ".tif"]:
            text = await parse_image_async(str(file_path_obj), params=params)
            result = f"{text}"

        elif file_ext in [".html", ".htm"]:
            async with aiofiles.open(file_path_obj, encoding="utf-8") as f:
                content = await f.read()
            text = await asyncio.to_thread(md_convert, content, heading_style="ATX")
            result = f"{text}"

        elif file_ext == ".csv":
            result = await asyncio.to_thread(_convert_csv_to_markdown, file_path_obj)

        elif file_ext in [".xls", ".xlsx"]:
            result = await asyncio.to_thread(_convert_with_docling, file_path_obj, params=params)

        elif file_ext == ".json":
            import json

            async with aiofiles.open(file_path_obj, encoding="utf-8") as f:
                content = await f.read()
            data = json.loads(content)
            json_str = json.dumps(data, ensure_ascii=False, indent=2)
            result = f"```json\n{json_str}\n```"

        elif file_ext == ".zip":
            image_bucket, image_prefix = _resolve_image_storage_params(params)
            zip_result = await _process_zip_file(
                str(file_path_obj),
                image_bucket=image_bucket,
                image_prefix=image_prefix,
            )

            artifacts = {
                "zip_images_info": zip_result["images_info"],
                "zip_content_hash": zip_result["content_hash"],
                "zip_image_bucket": image_bucket,
                "zip_image_prefix": image_prefix,
            }

            result = zip_result["markdown_content"]

        else:
            raise ValueError(f"Unsupported file type: {file_ext}")

    except Exception:
        if is_minio_url(file_path) and os.path.exists(actual_file_path):
            try:
                os.unlink(actual_file_path)
                logger.debug(f"Cleaned up temp file: {actual_file_path}")
            except Exception as cleanup_e:  # noqa: BLE001
                logger.warning(f"Failed to clean up temp file {actual_file_path}: {cleanup_e}")
        raise

    finally:
        if is_minio_url(file_path) and os.path.exists(actual_file_path):
            try:
                os.unlink(actual_file_path)
                logger.debug(f"Cleaned up temp file: {actual_file_path}")
            except Exception as e:  # noqa: BLE001
                logger.warning(f"Failed to clean up temp file {actual_file_path}: {e}")

    return result, file_ext, artifacts


async def parse_source_to_markdown(source: str, params: dict | None = None) -> MarkdownParseResult:
    """unified entrance: Parse files into Markdown (URL parsing is deprecated)."""
    markdown, file_ext, artifacts = await _process_file_to_markdown_core(source, params=params)
    return MarkdownParseResult(
        markdown=markdown,
        file_ext=file_ext,
        artifacts=artifacts,
    )


class Parser:
    """Lightweight facade for converting file sources to markdown."""

    @staticmethod
    async def aparse(source: str, params: dict | None = None) -> str:
        """Asynchronously parse source content and return markdown text."""
        parsed = await parse_source_to_markdown(source=source, params=params)
        return parsed.markdown

    @classmethod
    def parse(cls, source: str, params: dict | None = None) -> str:
        """Synchronously parse source content and return markdown text."""
        try:
            asyncio.get_running_loop()
        except RuntimeError:
            return asyncio.run(cls.aparse(source=source, params=params))

        raise RuntimeError("Hiện đang trong bối cảnh bất đồng bộ, vui lòng sử dụng `await Parser.aparse(...)`")
